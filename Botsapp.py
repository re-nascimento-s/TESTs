from flask import Flask, render_template, request, redirect, url_for
import openpyxl
from urllib.parse import quote
import webbrowser
import pyautogui
import pyperclip
import threading
import os

app = Flask(__name__)

# Variáveis globais
workbook = None
is_running = False
stop_event = threading.Event()

# Função para carregar o banco de dados
def carregar_banco(file_path):
    global workbook
    if file_path:
        workbook = openpyxl.load_workbook(file_path)
    return os.path.basename(file_path) if workbook else None

# Função para disparar mensagens
def disparar_mensagem(analista_nome, mensagem1, mensagem2, mensagem3):
    global is_running
    if not workbook or not analista_nome:
        return "Erro: Falha ao carregar o banco de dados ou nome do analista não fornecido"

    pagina_clientes = workbook['Planilha1']
    webbrowser.open('https://web.whatsapp.com/')
    sleep(8)  # Espera o WhatsApp Web carregar

    is_running = True  # Inicia o envio de mensagens
    for linha in pagina_clientes.iter_rows(min_row=2):
        if stop_event.is_set():
            break  # Interrompe o loop imediatamente

        nome = linha[0].value
        telefone = linha[1].value

        # Mensagens personalizadas com placeholders
        msg1 = mensagem1.format(nome=nome, analista=analista_nome)
        msg2 = mensagem2.format(nome=nome, analista=analista_nome)
        msg3 = mensagem3.format(nome=nome, analista=analista_nome)

        try:
            link_mensagem_whatsapp = f'https://web.whatsapp.com/send?phone={telefone}&text={quote(msg1)}'
            webbrowser.open(link_mensagem_whatsapp)
            sleep(12)
            pyautogui.press('enter')
            sleep(2)

            if stop_event.is_set():
                break

            pyperclip.copy(msg2)
            pyautogui.hotkey('ctrl', 'v')
            sleep(2)
            pyautogui.press('enter')
            sleep(2)

            if stop_event.is_set():
                break

            pyperclip.copy(msg3)
            pyautogui.hotkey('ctrl', 'v')
            sleep(2)
            pyautogui.press('enter')
            sleep(2)

            pyautogui.hotkey('ctrl', 'w')
            sleep(2)

            if stop_event.is_set():
                break

        except Exception as e:
            print(f'Erro ao enviar mensagem para {nome}: {e}')

    is_running = False  # Finaliza o processo
    stop_event.clear()  # Limpa o evento de parada

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/carregar_arquivo', methods=['POST'])
def carregar_arquivo():
    file = request.files['file']
    if file:
        file_path = os.path.join('uploads', file.filename)
        file.save(file_path)
        nome_arquivo = carregar_banco(file_path)
        return render_template('index.html', nome_arquivo=nome_arquivo)
    return redirect(url_for('index'))

@app.route('/enviar_mensagens', methods=['POST'])
def enviar_mensagens():
    analista_nome = request.form['analista']
    mensagem1 = request.form['mensagem1']
    mensagem2 = request.form['mensagem2']
    mensagem3 = request.form['mensagem3']

    if not analista_nome or not workbook:
        return "Erro: Nome do analista ou arquivo não carregado"

    # Inicia a thread para disparar as mensagens
    threading.Thread(target=disparar_mensagem, args=(analista_nome, mensagem1, mensagem2, mensagem3), daemon=True).start()
    return render_template('index.html', is_running=True)

@app.route('/parar_mensagens', methods=['POST'])
def parar_mensagens():
    stop_event.set()  # Define o evento de parada
    return render_template('index.html', is_running=False)

if __name__ == '__main__':
    app.run(debug=True)
